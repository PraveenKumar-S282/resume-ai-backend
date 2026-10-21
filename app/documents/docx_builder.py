from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.schemas.documents import ResumeContent


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _bottom_border(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": color.lstrip("#")}.items():
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def _setup(cfg: dict) -> Document:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = cfg["docx_font"]
    normal.font.size = Pt(cfg["body_pt"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), cfg["docx_font"])
    normal.paragraph_format.space_after = Pt(0)
    return doc


def _run(p, text, cfg, bold=False, size=None, color=None, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = cfg["docx_font"]
    r.font.size = Pt(size or cfg["body_pt"])
    if color:
        r.font.color.rgb = _rgb(color)
    return r


def _heading(doc, title, cfg):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(cfg["space"] + 4)
    p.paragraph_format.space_after = Pt(cfg["space"] - 2)
    p.paragraph_format.keep_with_next = True
    _run(p, title, cfg, bold=True, size=cfg["heading_pt"], color=cfg["accent"])
    _bottom_border(p, cfg["accent"])


def _bullets(doc, items, cfg):
    for b in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        _run(p, b, cfg)


def _two_col(doc, left_runs, right, cfg):
    """Left text with right-aligned date using a right tab stop."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(cfg["space"] - 1)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), alignment=2)  # right
    for text, bold in left_runs:
        _run(p, text, cfg, bold=bold)
    if right:
        _run(p, "\t" + right, cfg)
    p.paragraph_format.keep_with_next = True


def build_resume_docx(c: ResumeContent, cfg: dict) -> bytes:
    doc = _setup(cfg)
    h = c.header
    align = WD_ALIGN_PARAGRAPH.CENTER if cfg["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT

    p = doc.add_paragraph()
    p.alignment = align
    _run(p, h.name or "Your name", cfg, bold=True, size=cfg["name_pt"], color=cfg["accent"])
    if h.headline:
        p = doc.add_paragraph()
        p.alignment = align
        _run(p, h.headline, cfg, size=cfg["body_pt"] + 1)
    contact = [x for x in [h.email, h.phone, h.location, *[(l.label or l.url) for l in h.links]] if x]
    if contact:
        p = doc.add_paragraph()
        p.alignment = align
        _run(p, "  |  ".join(contact), cfg, color="#475569")

    if c.summary:
        _heading(doc, "Summary", cfg)
        _run(doc.add_paragraph(), c.summary, cfg)
    if c.skills:
        _heading(doc, "Skills", cfg)
        for g in c.skills:
            p = doc.add_paragraph()
            _run(p, f"{g.category}: ", cfg, bold=True)
            _run(p, ", ".join(g.items), cfg)
    if c.experience:
        _heading(doc, "Experience", cfg)
        for e in c.experience:
            left = [(e.role, True)] + ([(f", {e.company}", False)] if e.company else []) + ([(f", {e.location}", False)] if e.location else [])
            _two_col(doc, left, " to ".join(x for x in (e.start, e.end) if x), cfg)
            _bullets(doc, e.bullets, cfg)
    if c.projects:
        _heading(doc, "Projects", cfg)
        for pr in c.projects:
            left = [(pr.title, True)] + ([(f" ({', '.join(pr.tech)})", False)] if pr.tech else [])
            _two_col(doc, left, None, cfg)
            _bullets(doc, pr.bullets, cfg)
    if c.education:
        _heading(doc, "Education", cfg)
        for ed in c.education:
            left = [(ed.institution, True)] + ([(f", {ed.degree}", False)] if ed.degree else []) + ([(f" ({ed.grade})", False)] if ed.grade else [])
            _two_col(doc, left, ed.dates, cfg)
    if c.certifications:
        _heading(doc, "Certifications", cfg)
        for ct in c.certifications:
            line = ct.name + (f", {ct.issuer}" if ct.issuer else "") + (f" ({ct.date})" if ct.date else "")
            _run(doc.add_paragraph(), line, cfg)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_cover_letter_docx(text: str, cfg: dict) -> bytes:
    doc = _setup(cfg)
    for para in [p for p in text.split("\n\n") if p.strip()]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        lines = para.split("\n")
        for i, line in enumerate(lines):
            _run(p, line, cfg, size=cfg["body_pt"] + 1)
            if i < len(lines) - 1:
                p.add_run().add_break()
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
